# -*- coding: utf-8 -*-
"""tin-insight DATA 块降级刷新第二波：本地 Excel / 缓存镜像 / 公开源。
在 refresh_data_fallback.py 已刷键之外，覆盖其余过时图表：
  - Mysteel 锡线下数据.xlsx（贸易商成交/升贴水云锡/进出口盈亏/加工费/冶炼利润/
    下游焊料/电子产品/马口铁/美国日本精锡进口/印尼出口印度）
  - 锡数据库.xlsx（国内精炼锡概况：精炼锡/再生锡月度产量）
  - 缓存镜像 数据缓存_v1.db（LME 0-3 升贴水/总持仓/美国精锡进口/集成电路/SOX）
  - 新浪日线 akshare（SN0/AG0/CU0/AU0 主连 + 全合约持仓加总）+ 中证1000 + SPX(.INX)
合并铁律：与 DATA 现有序列取并集，同日冲突以新源为准，只前进不回落。
用法：python scripts/refresh_data_excel.py   （可重复运行，幂等）
"""
import datetime
import json
import math
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from refresh_data_fallback import (  # noqa: E402
    COLORS, cache_series, continuous_chart, latest, merge, meta_obj,
    seasonal_chart, uncontinuous, unseasonal,
)

ROOT = Path(__file__).resolve().parent.parent
MYSTEEL = r"D:\拷贝文件\E\永安\周报数据更新\Mysteel锡线下数据.xlsx"
SMM_DB = r"D:\拷贝文件\E\永安\周报数据更新\【永安期货研究中心】锡数据库.xlsx"

REPORT = []


def iso(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    return None


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
        if math.isfinite(f):
            return f
    return None


def sheet_series(ws, date_col, val_col, min_row=1, scale=1.0, nd=6):
    """抽 (date, value) 序列，升序返回。列号为 1-based。"""
    out = []
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if len(row) < max(date_col, val_col):
            continue
        d = iso(row[date_col - 1])
        v = num(row[val_col - 1])
        if d and v is not None:
            out.append((d, round(v * scale, nd)))
    out.sort()
    return out


def last_date_of_series(items):
    return items[-1][0] if items else None


def chart_last(chart):
    """chart 现有最后有效日期（季节性拼回完整日期）。"""
    labels = chart.get("labels") or []
    best = None
    for ds in chart.get("datasets") or []:
        y = ds.get("label", "")
        for lab, v in zip(labels, ds.get("data") or []):
            if v is None:
                continue
            d = f"{y}-{lab}" if y.isdigit() else lab
            if best is None or d > best:
                best = d
    return best


def guard(chart_key, merged, old_last):
    """只前进不回落：合并后 last 不得早于 DATA 现有 last。"""
    new_last = last_date_of_series(merged)
    if old_last and (not new_last or new_last < old_last):
        REPORT.append(f"{chart_key}: 新源回落({new_last}<{old_last})，放弃")
        return None
    return new_last


def main():
    import openpyxl

    text = (ROOT / "index.html").read_text(encoding="utf-8")
    prefix, suffix = "const DATA=", ";\nconst STATIC_HOST="
    start = text.index(prefix) + len(prefix)
    end = text.index(suffix, start)
    data = json.loads(text[start:end])
    charts = data.setdefault("charts", {})
    latest_data = data.setdefault("latest", {})
    source_meta = data.setdefault("sourceMeta", {})

    # 快照每个 chart 现有 dataset 样式（重建后按位置恢复，保持配色/线宽不变）
    old_styles = {
        k: [{kk: ds.get(kk) for kk in ("borderColor", "backgroundColor", "borderWidth")}
            for ds in (c.get("datasets") or [])]
        for k, c in charts.items()
    }

    wb_m = openpyxl.load_workbook(MYSTEEL, read_only=True, data_only=True)
    wb_s = openpyxl.load_workbook(SMM_DB, read_only=True, data_only=True)

    def refresh_seasonal(key, new_series, latest_key=None):
        """季节性 chart 合并刷新。new_series: [(date, value)]。返回合并后最后日期。"""
        if not new_series:
            REPORT.append(f"{key}: 新源为空，保留现状")
            return None
        old_last = chart_last(charts.get(key, {}))
        old = unseasonal(charts.get(key, {}))
        merged = merge(old, dict(new_series))
        nl = guard(key, merged, old_last)
        if nl is None:
            return None
        charts[key] = seasonal_chart(merged)
        if latest_key:
            latest_data[latest_key] = latest(merged)
        REPORT.append(f"{key}: {old_last} -> {nl} (+{len(merged) - len(old)} 点)")
        return nl

    # meta 参数打包: (key, id, name, unit, org, source_note)
    def set_meta(key, ind_id, name, unit, org, source, last_d):
        source_meta[key] = meta_obj(ind_id, name, unit, org, last_d, source)

    # ============ 1. overview_shfe ← 新浪 SN0 收盘 ============
    sina = {}
    try:
        import akshare as ak
        for sym in ["SN0", "AG0", "CU0", "AU0"]:
            df = ak.futures_zh_daily_sina(symbol=sym)
            sina[sym] = [(str(d)[:10], float(c)) for d, c in zip(df["date"], df["close"])]
            time.sleep(1.0)
        df = ak.stock_zh_index_daily(symbol="sh000852")
        sina["CSI1000"] = [(str(d)[:10], float(c)) for d, c in zip(df["date"], df["close"])]
        time.sleep(1.0)
        df = ak.stock_us_daily(symbol=".INX")
        sina["SPX"] = [(str(d)[:10], float(c)) for d, c in zip(df["date"], df["close"])]
        REPORT.append(f"新浪/指数源: { {k: v[-1][0] for k, v in sina.items()} }")
    except Exception as e:
        REPORT.append(f"新浪/指数源抓取失败: {e!r}")

    if sina.get("SN0"):
        # 口径校验：DATA latest.shfe 应与 SN0 同口径（收盘）
        lv = latest_data.get("shfe")
        if lv and lv[0]:
            chk = dict(sina["SN0"]).get(lv[0])
            if chk is not None and abs(chk - lv[1]) > 1:
                REPORT.append(f"overview_shfe: 口径校验不一致 DATA({lv[0]}={lv[1]}) vs SN0({chk})，仍按收盘合并")
        nl = refresh_seasonal("overview_shfe", sina["SN0"], latest_key="shfe")
        if nl:
            set_meta("shfePrice", "SN0", "沪锡主连收盘价", "元/吨", "上期所/新浪", "新浪日线", nl)

    # ============ 2. overview_trade ← Mysteel 贸易商日度交易量 ============
    ws = wb_m["贸易商日度交易量"]
    peer = sheet_series(ws, 1, 2, min_row=3)
    down = sheet_series(ws, 1, 3, min_row=3)
    if peer or down:
        key = "overview_trade"
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        name_p, name_d = "锡锭同业贸易商成交合计", "锡锭下游成交合计"
        m_p = merge(old.get(name_p, {}), dict(peer))
        m_d = merge(old.get(name_d, {}), dict(down))
        nl = max(last_date_of_series(m_p) or "", last_date_of_series(m_d) or "")
        if old_last and nl < old_last:
            REPORT.append(f"{key}: 回落，放弃")
        else:
            charts[key] = continuous_chart({name_p: m_p, name_d: m_d}, limit=None)
            latest_data["trade_peer"] = latest(m_p)
            latest_data["trade_downstream"] = latest(m_d)
            set_meta("trade", "mysteel-trade", "锡锭贸易商日度成交", "吨", "Mysteel",
                     "Mysteel锡线下数据.xlsx", nl)
            REPORT.append(f"{key}: {old_last} -> {nl}")

    # ============ 3. premium ← Mysteel 升贴水（仅云锡列口径一致，其余保留）============
    ws = wb_m["升贴水"]
    yunxi = sheet_series(ws, 1, 2, min_row=6)  # 云锡 升贴水
    if yunxi:
        key = "premium"
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        # 口径校验：衔接段（尾部 5 个共同日）应一致；历史修订不改写，只做纯延伸
        ovl = old.get("云锡", {})
        old_yx_last = sorted(ovl.items())[-1][0] if ovl else None
        common = sorted(d for d, _ in yunxi if d in ovl)[-5:]
        same = all(abs(ovl[d] - dict(yunxi)[d]) < 1e-6 for d in common)
        if ovl and (not common or not same):
            REPORT.append(f"{key}: 云锡衔接段口径不一致，保留现状")
        else:
            ext = [(d, v) for d, v in yunxi if old_yx_last is None or d > old_yx_last]
            order = [ds.get("label") for ds in charts[key]["datasets"]]
            merged_map = {}
            for name in order:
                s = old.get(name, {})
                if name == "云锡" and ext:
                    merged_map[name] = merge(s, dict(ext))  # ext 全在 old_last 之后，不触历史
                else:
                    merged_map[name] = sorted(s.items())
            new_last = max((last_date_of_series(v) for v in merged_map.values() if v), default=None)
            if old_last and new_last and new_last < old_last:
                REPORT.append(f"{key}: 回落，放弃")
            else:
                limit = len(charts[key].get("labels") or []) or None
                charts[key] = continuous_chart(merged_map, limit=limit)
                set_meta("premium", "mysteel-premium", "锡锭升贴水（云锡）", "元/吨", "Mysteel",
                         "Mysteel锡线下数据.xlsx", new_last)
                REPORT.append(f"{key}(云锡): {old_last} -> {new_last}")

    # ============ 4. lme03 ← 缓存 a10099105（口径已校验=DATA 8/14 值）============
    ser, meta = cache_series("a10099105")
    if ser:
        refresh_seasonal("lme03", ser, latest_key="lme03")
        m = merge(unseasonal(charts["lme03"]), {})
        set_meta("lme03", "a10099105", meta[0], meta[1], meta[2], "缓存镜像(zhiji/SMM)", m[-1][0])

    # ============ 5. import_profit ← 缓存 ID01590826 + Mysteel 进出口盈亏 J 列 ============
    ws = wb_m["进出口盈亏"]
    mys_profit = sheet_series(ws, 1, 10, min_row=2, nd=2)  # 现货进口盈亏
    ser, meta = cache_series("ID01590826")
    comb = merge(dict(ser), dict(mys_profit))
    if comb:
        refresh_seasonal("import_profit", comb, latest_key="import_profit")
        m = merge(unseasonal(charts["import_profit"]), {})
        set_meta("importProfit", "ID01590826+mysteel", "锡锭进口利润", "元/吨", "Mysteel/SMM",
                 "Mysteel锡线下数据.xlsx+缓存镜像", m[-1][0])

    # ============ 6. sn_total_oi ← 缓存 FU00017555 + 新浪全合约持仓加总 ============
    ser, meta = cache_series("FU00017555")
    sina_oi = []
    try:
        import akshare as ak
        import pandas as pd
        holds = {}
        for y in (25, 26, 27):
            for mo in range(1, 13):
                sym = f"SN{y}{mo:02d}"
                try:
                    df = ak.futures_zh_daily_sina(symbol=sym)
                except Exception:
                    continue
                if df is None or not len(df):
                    continue
                for d, h in zip(df["date"], df["hold"]):
                    v = num(h)
                    if v:
                        holds[str(d)[:10]] = holds.get(str(d)[:10], 0) + v
                time.sleep(1.0)
        sina_oi = sorted(holds.items())
    except Exception as e:
        REPORT.append(f"sn_total_oi 新浪加总失败: {e!r}")
    if sina_oi and ser:
        # 一致性校验：缓存最后 5 个共同日相对偏差
        cache_map = dict(ser)
        common = [(d, v, cache_map[d]) for d, v in sina_oi if d in cache_map][-5:]
        if common:
            dev = max(abs(a - b) / b for _, a, b in common if b)
            REPORT.append(f"sn_total_oi 新浪加总 vs 缓存 最大相对偏差={dev:.4f}")
            if dev > 0.02:
                REPORT.append("sn_total_oi: 偏差>2%，只用缓存")
                sina_oi = []
    comb = merge(dict(ser), dict(sina_oi[-30:]))  # 新浪只补尾部
    if comb:
        key = "sn_total_oi"
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        name = "沪锡总持仓量"
        m = merge(old.get(name, {}), dict(comb))
        nl = last_date_of_series(m)
        if old_last and nl < old_last:
            REPORT.append(f"{key}: 回落，放弃")
        else:
            charts[key] = continuous_chart({name: m}, limit=500)
            latest_data["sn_total_oi"] = latest(m)
            set_meta("snTotalOi", "FU00017555+sina", meta[0], meta[1], "上期所",
                     "缓存镜像(zhiji)+新浪日线加总", nl)
            REPORT.append(f"{key}: {old_last} -> {nl}")

    # ============ 7. corr_tin / corr_metals ← 60 日滚动相关重算 ============
    try:
        import pandas as pd
        sox, _ = cache_series("GM00194406")
        if sina.get("SN0") and sina.get("SPX"):
            def ret_series(items):
                s = pd.Series({d: v for d, v in items}).sort_index()
                return s
            tin = ret_series(sina["SN0"])
            spx = ret_series(sina["SPX"]).reindex(tin.index).ffill()
            def roll_corr(other, win=60):
                o = other.reindex(tin.index).ffill()
                r = pd.concat([tin.pct_change(), o.pct_change()], axis=1).dropna()
                c = r.iloc[:, 0].rolling(win).corr(r.iloc[:, 1]).dropna()
                return [(d, round(float(v), 4)) for d, v in c.items()]
            pairs_tin = [("锡 vs SPX", roll_corr(spx))]
            if sox:
                pairs_tin.append(("锡 vs SOX", roll_corr(ret_series(sox))))
            for key, pairs in [("corr_tin", pairs_tin)]:
                old_last = chart_last(charts.get(key, {}))
                series_map = {}
                for base, serie in pairs:
                    if not serie:
                        continue
                    now = serie[-1][1]
                    series_map[f"{base} (now={now:.3f})"] = serie
                if series_map:
                    nl = max(v[-1][0] for v in series_map.values())
                    if old_last and nl < old_last:
                        REPORT.append(f"{key}: 回落，放弃")
                    else:
                        charts[key] = continuous_chart(series_map, limit=180)
                        set_meta(key, "computed", "60日滚动相关（主连日收益率）", "-", "计算",
                                 "新浪日线+stooq/sina SPX+缓存SOX", nl)
                        REPORT.append(f"{key}: {old_last} -> {nl} nows={[v[-1][1] for v in series_map.values()]}")
            # corr_metals: 锡/银/铜/金 vs SPX
            pairs_m = [("锡 vs SPX", pairs_tin[0][1])]
            for base, sym in [("银 vs SPX", "AG0"), ("铜 vs SPX", "CU0"), ("金 vs SPX", "AU0")]:
                if sina.get(sym):
                    pairs_m.append((base, roll_corr(ret_series(sina[sym]))))
            key = "corr_metals"
            old_last = chart_last(charts.get(key, {}))
            series_map = {f"{b} (now={s[-1][1]:.3f})": s for b, s in pairs_m if s}
            if series_map:
                nl = max(v[-1][0] for v in series_map.values())
                if old_last and nl < old_last:
                    REPORT.append(f"{key}: 回落，放弃")
                else:
                    charts[key] = continuous_chart(series_map, limit=180)
                    set_meta(key, "computed", "60日滚动相关（主连日收益率）", "-", "计算",
                             "新浪日线+sina SPX", nl)
                    REPORT.append(f"{key}: {old_last} -> {nl}")
    except Exception as e:
        REPORT.append(f"corr 重算失败: {e!r}")

    # ============ 8. ratio 三图 ← 新浪收盘比值 ============
    ratio_specs = [
        ("ratio_tin_silver", "AG0", "沪锡 / 沪银", "tin_silver_ratio", "ratioTinSilver"),
        ("ratio_tin_copper", "CU0", "沪锡 / 沪铜", "tin_copper_ratio", "ratioTinCopper"),
        ("ratio_tin_csi1000", "CSI1000", "沪锡 / 中证1000", "tin_csi1000_ratio", "ratioTinCsi1000"),
    ]
    for key, sym, label, latest_key, meta_key in ratio_specs:
        if not (sina.get("SN0") and sina.get(sym)):
            REPORT.append(f"{key}: 缺源，保留现状")
            continue
        other = dict(sina[sym])
        serie = [(d, round(t / other[d], 6)) for d, t in sina["SN0"] if d in other and other[d]]
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        m = merge(old.get(label, {}), dict(serie))
        nl = last_date_of_series(m)
        if old_last and nl < old_last:
            REPORT.append(f"{key}: 回落，放弃")
            continue
        charts[key] = continuous_chart({label: m}, limit=600)
        latest_data[latest_key] = latest(m)
        set_meta(meta_key, "computed", label, "比值", "计算", "新浪日线", nl)
        REPORT.append(f"{key}: {old_last} -> {nl}")

    # ============ 9. refined / recycled ← 锡数据库 国内精炼锡概况 K/L、AE/AF ============
    ws = wb_s["国内精炼锡概况"]
    ref = sheet_series(ws, 11, 12)
    rec = sheet_series(ws, 31, 32)
    refresh_seasonal("refined", ref)
    if ref:
        m = merge(unseasonal(charts["refined"]), {})
        set_meta("refinedProduction", "a10003083", "国内精炼锡产量（SMM)", "吨", "SMM",
                 "锡数据库.xlsx", m[-1][0])
    refresh_seasonal("recycled", rec)
    if rec:
        m = merge(unseasonal(charts["recycled"]), {})
        set_meta("recycledProduction", "a10099502", "国内再生锡产量（SMM)", "吨", "SMM",
                 "锡数据库.xlsx", m[-1][0])

    # ============ 10. ic ← Mysteel 电子产品 集成电路 ×10000 + 缓存 ============
    ws = wb_m["电子产品 "]
    ic_mys = sheet_series(ws, 1, 3, min_row=6, scale=10000)
    ser, meta = cache_series("CM0000017738")
    ic_cache = [(d, round(v * 10000)) for d, v in ser]
    comb = merge(dict(ic_cache), dict(ic_mys))
    if comb:
        refresh_seasonal("ic", comb)
        m = merge(unseasonal(charts["ic"]), {})
        set_meta("icProduction", "CM0000017738", "中国集成电路产量", "万块", "统计局/Mysteel",
                 "Mysteel锡线下数据.xlsx+缓存镜像", m[-1][0])

    # ============ 11. tinplate ← Mysteel 马口铁产量 宽表（E=月份 F..=年份）============
    ws = wb_m["马口铁产量"]
    rows = list(ws.iter_rows(min_row=1, max_row=13, max_col=12, values_only=True))
    year_cols = {}
    for j, v in enumerate(rows[0]):
        s = str(v or "")
        if s.endswith("年") and s[:-1].isdigit():
            year_cols[int(s[:-1])] = j
    tp = []
    for r in rows[1:]:
        ms = str(r[4] or "")
        if ms.endswith("月") and ms[:-1].isdigit():
            mo = int(ms[:-1])
            for y, j in year_cols.items():
                v = num(r[j]) if len(r) > j else None
                if v is not None:
                    tp.append((f"{y}-{mo:02d}-01", v))
    tp.sort()
    if tp:
        refresh_seasonal("tinplate", tp)
        m = merge(unseasonal(charts["tinplate"]), {})
        set_meta("tinplate", "mysteel-tinplate", "马口铁产量", "万吨", "Mysteel",
                 "Mysteel锡线下数据.xlsx", m[-1][0])

    # ============ 12. solder ← Mysteel 下游焊料厂产量 ============
    ws = wb_m["下游焊料厂产量"]
    so = sheet_series(ws, 1, 2, min_row=2)
    if so:
        refresh_seasonal("solder", so)
        m = merge(unseasonal(charts["solder"]), {})
        set_meta("solder", "mysteel-solder", "下游焊料厂产量", "吨", "Mysteel",
                 "Mysteel锡线下数据.xlsx", m[-1][0])

    # ============ 13. us_import ← Mysteel 美国精锡进口 + 缓存 ID01975601(kg→t) ============
    ws = wb_m["美国精锡进口"]
    us_mys = sheet_series(ws, 1, 2, min_row=2)
    ser, meta = cache_series("ID01975601")
    us_cache = [(d, round(v / 1000, 4)) for d, v in ser]
    comb = merge(dict(us_mys), dict(us_cache))
    if comb:
        refresh_seasonal("us_import", comb)
        m = merge(unseasonal(charts["us_import"]), {})
        set_meta("usImport", "ID01975601", meta[0], "吨", meta[2],
                 "缓存镜像(zhiji)+Mysteel锡线下数据.xlsx", m[-1][0])

    # ============ 14. jp_import ← Mysteel 日本精炼锡进口 B 列(kg→t) ============
    ws = wb_m["日本精炼锡进口数量"]
    jp = sheet_series(ws, 1, 2, min_row=6, scale=0.001, nd=4)
    if jp:
        refresh_seasonal("jp_import", jp)
        m = merge(unseasonal(charts["jp_import"]), {})
        set_meta("jpImport", "mysteel-jp", "日本精炼锡进口量", "吨", "日本海关/Mysteel",
                 "Mysteel锡线下数据.xlsx", m[-1][0])

    # ============ 15. kr_import ← Mysteel 韩国 B 列（持平校验用，无新增则不动）============
    ws = wb_m["韩国精锡及锡制品进出口"]
    kr = sheet_series(ws, 1, 2, min_row=6)
    if kr:
        refresh_seasonal("kr_import", kr)

    # ============ 16. indo_export_india ← Mysteel 印尼精锡出口 F 列（至印度）============
    ws = wb_m["印尼精锡出口"]
    ind = sheet_series(ws, 1, 6, min_row=20)
    if ind:
        refresh_seasonal("indo_export_india", ind)
        m = merge(unseasonal(charts["indo_export_india"]), {})
        set_meta("indoExportIndia", "mysteel-ido-india", "印尼精炼锡出口（至印度）", "吨", "Mysteel",
                 "Mysteel锡线下数据.xlsx", m[-1][0])

    # ============ 17. cost_tc / cost_profit 不回落合并（Mysteel 更鲜）============
    ws = wb_m["锡精矿加工费统计"]
    tc_y = sheet_series(ws, 1, 2, min_row=6)
    tc_j = sheet_series(ws, 1, 3, min_row=6)
    if tc_y or tc_j:
        key = "cost_tc"
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        name_y, name_j = "云南 40% 锡精矿加工费", "江西 60% 锡精矿加工费"
        m_y = merge(old.get(name_y, {}), dict(tc_y))
        m_j = merge(old.get(name_j, {}), dict(tc_j))
        nl = max(last_date_of_series(m_y) or "", last_date_of_series(m_j) or "")
        if old_last and nl < old_last:
            REPORT.append(f"{key}: 回落，放弃")
        else:
            charts[key] = continuous_chart({name_y: m_y, name_j: m_j}, limit=260)
            latest_data["tc_yunnan"] = latest(m_y)
            latest_data["tc_jiangxi"] = latest(m_j)
            set_meta("tcYunnan", "mysteel-tc", "锡精矿TC（云南40%）", "元/吨", "Mysteel",
                     "Mysteel锡线下数据.xlsx", nl)
            set_meta("tcJiangxi", "mysteel-tc", "锡精矿TC（江西60%）", "元/吨", "Mysteel",
                     "Mysteel锡线下数据.xlsx", nl)
            REPORT.append(f"{key}: {old_last} -> {nl}")
    ws = wb_m["锡冶炼利润"]
    prof = sheet_series(ws, 1, 10, nd=2)
    if prof:
        key = "cost_profit"
        old = uncontinuous(charts.get(key, {}))
        old_last = chart_last(charts.get(key, {}))
        name = "冶炼锡毛利（日度）"
        m_p = merge(old.get(name, {}), dict(prof))
        nl = last_date_of_series(m_p)
        if old_last and nl < old_last:
            REPORT.append(f"{key}: 回落，放弃")
        else:
            charts[key] = continuous_chart({name: m_p}, limit=260)
            latest_data["smelting_profit"] = latest(m_p)
            set_meta("smeltingProfit", "mysteel-profit", "锡冶炼利润", "元/吨", "Mysteel",
                     "Mysteel锡线下数据.xlsx", nl)
            REPORT.append(f"{key}: {old_last} -> {nl}")

    wb_m.close()
    wb_s.close()

    # 恢复所有重建 chart 的 dataset 样式（按位置；corr 标签文字含 now= 会变，样式不变）
    for k, c in charts.items():
        styles = old_styles.get(k) or []
        for ds, st in zip(c.get("datasets") or [], styles):
            for kk, vv in st.items():
                if vv is None:
                    ds.pop(kk, None)
                else:
                    ds[kk] = vv

    out = text[:start] + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + text[end:]
    (ROOT / "index.html").write_text(out, encoding="utf-8", newline="")
    print("\n".join(REPORT))
    print(f"index.html 已更新 ({len(out)/1e3:.0f} KB)")


if __name__ == "__main__":
    sys.exit(main())
